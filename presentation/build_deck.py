"""Generates presentation/ConvoScore-Presentation.pptx.

Build-time tooling only (python-pptx) - not an application runtime
dependency, not referenced anywhere under app/. Run with:
    .venv/Scripts/python.exe presentation/build_deck.py

Editing: this file *is* the editable source for the deck - change the
slide functions below and re-run to regenerate the .pptx.

Design system: 10 slides, one idea each, <=4 top-level points on most
slides. Deeper technical detail (claim-fencing internals, exact IAM
actions, dirty-tree checks, refusal handling, retry implementation,
health-gauge internals) lives in speaker notes for Q&A, not on slides.
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
SCREENSHOTS = HERE / "screenshots"
OUT = HERE / "ConvoScore-Presentation.pptx"

# --- palette -------------------------------------------------------------
NAVY = RGBColor(0x14, 0x25, 0x3A)
NAVY_2 = RGBColor(0x1C, 0x33, 0x4D)
SLATE = RGBColor(0x4A, 0x5A, 0x6A)
TEAL = RGBColor(0x16, 0x8A, 0x86)
TEAL_DARK = RGBColor(0x0F, 0x66, 0x64)
AMBER = RGBColor(0xC9, 0x7A, 0x1A)
DANGER = RGBColor(0x9E, 0x35, 0x33)
LIGHT_BG = RGBColor(0xF7, 0xF8, 0xFA)
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)
CARD_LINE = RGBColor(0xE1, 0xE5, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x24, 0x2B)
MUTED = RGBColor(0x62, 0x6E, 0x79)
FAINT = RGBColor(0xB7, 0xC2, 0xCC)

FONT = "Segoe UI"
FONT_SEMI = "Segoe UI Semibold"

TOTAL_SLIDES = 10
SLIDE_W = 13.333
MARGIN = 0.7
CONTENT_R = SLIDE_W - MARGIN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# --- low-level helpers -----------------------------------------------------

def add_slide():
    return prs.slides.add_slide(BLANK)


def set_background(slide, color=LIGHT_BG):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _font(run, size, color, bold=False, name=FONT, italic=False, spc=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    if spc is not None:
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(spc * 100)))


def no_shadow(shape):
    shape.shadow.inherit = False


def rounded(shape, radius=0.06):
    try:
        shape.adjustments[0] = radius
    except (IndexError, AttributeError):
        pass


# --- shared slide chrome -----------------------------------------------

def kicker_title(slide, kicker, title, accent=TEAL, title_color=NAVY):
    k = slide.shapes.add_textbox(Inches(MARGIN), Inches(0.42), Inches(10.5), Inches(0.32))
    tf = k.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = kicker.upper()
    _font(r, 12.5, accent, bold=True, spc=1.6)

    t = slide.shapes.add_textbox(Inches(MARGIN - 0.03), Inches(0.72), Inches(11.6), Inches(0.7))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    title_size = 27 if len(title) <= 55 else 22
    _font(r, title_size, title_color, bold=True)

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(1.48), Inches(1.15), Pt(3))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()
    no_shadow(line)


def footer(slide, section, index, total=TOTAL_SLIDES):
    sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(6.86), Inches(CONTENT_R - MARGIN), Pt(0.75))
    sep.fill.solid()
    sep.fill.fore_color.rgb = CARD_LINE
    sep.line.fill.background()
    no_shadow(sep)

    l = slide.shapes.add_textbox(Inches(MARGIN), Inches(6.98), Inches(8.0), Inches(0.32))
    p = l.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"ConvoScore  ·  {section}"
    _font(r, 10, MUTED)

    rgt = slide.shapes.add_textbox(Inches(CONTENT_R - 1.3), Inches(6.98), Inches(1.3), Inches(0.32))
    p2 = rgt.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{index:02d} / {total}"
    _font(r2, 10, FAINT)


def bullets(slide, items, left, top, width, size=17, gap=13, color=INK,
            marker_color=TEAL, lead_color=NAVY, line_spacing=1.1, height=4.6):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        m = p.add_run()
        m.text = "▸  "
        _font(m, size, marker_color, bold=True)
        if " — " in item:
            lead, rest = item.split(" — ", 1)
            r1 = p.add_run()
            r1.text = lead
            _font(r1, size, lead_color, bold=True)
            r2 = p.add_run()
            r2.text = " — " + rest
            _font(r2, size, color)
        else:
            r1 = p.add_run()
            r1.text = item
            _font(r1, size, color)
    return box


def two_col_bullets(slide, left_items, right_items, top=1.95, col_w=5.55, gap_w=0.5,
                     size=15, height=4.4):
    lx = MARGIN
    rx = MARGIN + col_w + gap_w
    bullets(slide, left_items, lx, top, col_w, size=size, height=height)
    bullets(slide, right_items, rx, top, col_w, size=size, height=height)


def panel(slide, left, top, width, height, accent=TEAL):
    p = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    p.fill.solid()
    p.fill.fore_color.rgb = CARD_BG
    p.line.color.rgb = CARD_LINE
    p.line.width = Pt(0.75)
    no_shadow(p)
    rounded(p, 0.035)
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.16), Inches(top), Inches(width - 0.32), Pt(3.2))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    no_shadow(strip)
    return p


def panel_header(slide, text, left, top, width, color=NAVY, size=16):
    t = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.4))
    p = t.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    _font(r, size, color, bold=True)


def numbered_step(slide, n, text, left, top, width, accent=TEAL, size=13.5, d=0.36):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = accent
    c.line.fill.background()
    no_shadow(c)
    tf = c.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    _font(r, 13, WHITE, bold=True)

    tb = slide.shapes.add_textbox(Inches(left + d + 0.16), Inches(top - 0.09), Inches(width - d - 0.16), Inches(0.85))
    tf2 = tb.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.line_spacing = 1.05
    r2 = p2.add_run()
    r2.text = text
    _font(r2, size, INK)
    return tb


def add_picture_fit(slide, path, left, top, max_w, max_h, border=True):
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(Inches(max_w) / iw, Inches(max_h) / ih)
    w, h = Emu(int(iw * ratio)), Emu(int(ih * ratio))
    l = Inches(left) + (Inches(max_w) - w) / 2
    t = Inches(top) + (Inches(max_h) - h) / 2
    pic = slide.shapes.add_picture(str(path), l, t, width=w, height=h)
    if border:
        pic.line.color.rgb = CARD_LINE
        pic.line.width = Pt(1)
    return pic


# --- diagram primitives --------------------------------------------------

def dbox(slide, text, left, top, w, h, fill=TEAL, text_color=WHITE, size=12.5, subtitle=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    no_shadow(shp)
    rounded(shp, 0.12)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    _font(r, size, text_color, bold=True)
    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtitle
        _font(r2, size - 3.5, text_color)
    return shp


def seg(slide, x1, y1, x2, y2, color=SLATE, width=1.75, dashed=False, head=True):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    if dashed:
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(d)
    if head:
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return conn


def dlabel(slide, x, y, text, color=MUTED, size=10.5, align=PP_ALIGN.CENTER, width=2.2):
    lb = slide.shapes.add_textbox(Inches(x - width / 2), Inches(y), Inches(width), Inches(0.3))
    tf = lb.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _font(r, size, color)
    return lb


def elbow(slide, x1, y1, lane_y, x2, y2, color=AMBER, width=1.6, dashed=False):
    seg(slide, x1, y1, x1, lane_y, color=color, width=width, dashed=dashed, head=False)
    seg(slide, x1, lane_y, x2, lane_y, color=color, width=width, dashed=dashed, head=False)
    seg(slide, x2, lane_y, x2, y2, color=color, width=width, dashed=dashed, head=True)


def legend(slide, items, left, top, item_w=3.55):
    x = left
    for color, text, dashed in items:
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(top + 0.09), Inches(x + 0.32), Inches(top + 0.09))
        ln.line.color.rgb = color
        ln.line.width = Pt(2.0)
        if dashed:
            d = ln.line._get_or_add_ln().makeelement(qn("a:prstDash"), {"val": "dash"})
            ln.line._get_or_add_ln().append(d)
        lb = slide.shapes.add_textbox(Inches(x + 0.42), Inches(top - 0.06), Inches(item_w - 0.42), Inches(0.3))
        p = lb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = text
        _font(r, 10.5, MUTED)
        x += item_w


# ========================================================================
# Slide 1 - Title
# ========================================================================
s = add_slide()
set_background(s, NAVY)

band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
band.fill.solid()
band.fill.fore_color.rgb = NAVY
band.line.fill.background()
no_shadow(band)
accent_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Pt(6), Inches(7.5))
accent_bar.fill.solid()
accent_bar.fill.fore_color.rgb = TEAL
accent_bar.line.fill.background()
no_shadow(accent_bar)

k = s.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(10.5), Inches(0.35))
p = k.text_frame.paragraphs[0]
r = p.add_run()
r.text = "SOLUTION ENGINEERING / DEVOPS TAKE-HOME"
_font(r, 13, RGBColor(0x6F, 0xC7, 0xC2), bold=True, spc=2.0)

t = s.shapes.add_textbox(Inches(0.96), Inches(2.55), Inches(11.0), Inches(1.3))
p = t.text_frame.paragraphs[0]
r = p.add_run()
r.text = "ConvoScore"
_font(r, 56, WHITE, bold=True)

t2 = s.shapes.add_textbox(Inches(1.0), Inches(3.65), Inches(10.3), Inches(0.9))
tf = t2.text_frame
tf.word_wrap = True
p2 = tf.paragraphs[0]
r2 = p2.add_run()
r2.text = "Turning raw support conversations into a scored, reviewable, on-call-ready pipeline."
_font(r2, 19, RGBColor(0xC3, 0xCF, 0xD9))

t3 = s.shapes.add_textbox(Inches(1.0), Inches(6.55), Inches(10), Inches(0.4))
p3 = t3.text_frame.paragraphs[0]
r3 = p3.add_run()
r3.text = "Amir Weiser"
_font(r3, 13, RGBColor(0x9B, 0xAB, 0xB8), bold=True)
r3b = p3.add_run()
r3b.text = "   ·   Architecture, ingestion, scoring, deployment, and operations"
_font(r3b, 13, RGBColor(0x6E, 0x80, 0x8F))

add_notes(s, "One line: this is a small service that scores support conversations with an "
             "LLM and makes the results reviewable and observable - built to show how I'd "
             "design, deploy, and operate a real system, not just make a demo run once.")

# ========================================================================
# Slide 2 - What it does & why this architecture
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Overview", "What ConvoScore does, and why it's shaped this way")

panel(s, MARGIN, 1.9, 5.55, 4.35, accent=TEAL)
panel_header(s, "What it does", MARGIN + 0.3, 2.12, 5.0)
bullets(s, [
    "Ingests support conversations two ways — an API call, or a file dropped into storage",
    "Scores each one with an LLM — sentiment, a 0–1 risk score, a short rationale",
    "Stores results durably and makes them reviewable by a person — no separate frontend",
], MARGIN + 0.3, 2.65, 4.95, size=14.5, gap=16, height=3.4)

panel(s, MARGIN + 6.05, 1.9, 5.55, 4.35, accent=AMBER)
panel_header(s, "Why this architecture", MARGIN + 6.35, 2.12, 5.0)
bullets(s, [
    "Decouples ingestion from scoring — a slow or failed OpenAI call never blocks the caller",
    "One shared pipeline behind both entry points — one code path to trust and test",
    "Designed around two honest constraints — SQS is at-least-once, and an LLM call can't be made atomic with a DB commit",
], MARGIN + 6.35, 2.65, 4.95, size=14.5, gap=16, height=3.4, marker_color=AMBER)

footer(s, "Overview", 2)
add_notes(s, "This slide exists so the reviewer sees I read the brief as a set of constraints, "
             "not just a feature list. Additional scope notes if asked: kept deliberately "
             "right-sized — no CI, no service mesh, no ORM, no new database. The two "
             "constraints on the right (SQS at-least-once, LLM-call/DB-commit non-atomicity) "
             "shape almost every design choice in the rest of the deck.")

# ========================================================================
# Slide 3 - Architecture diagram
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Architecture", "One pipeline, two doors in")

API = dict(left=0.7, top=2.2, w=1.85, h=0.8)
DIRECT = dict(left=0.7, top=3.35, w=1.85, h=0.8)
S3 = dict(left=3.15, top=2.78, w=1.7, h=0.85)
SQS = dict(left=5.35, top=2.78, w=1.7, h=0.85)
WORKER = dict(left=7.55, top=2.78, w=1.55, h=0.85)
OPENAI = dict(left=7.55, top=4.05, w=1.55, h=0.75)
PG = dict(left=9.65, top=2.78, w=1.65, h=0.85)
REVIEW = dict(left=9.65, top=4.05, w=1.65, h=0.75)
DLQ = dict(left=5.35, top=4.05, w=1.7, h=0.75)


def ctr(b):
    return b["left"] + b["w"] / 2, b["top"] + b["h"] / 2


def edge(b, side):
    l, t, w, h = b["left"], b["top"], b["w"], b["h"]
    return {"l": (l, t + h / 2), "r": (l + w, t + h / 2), "t": (l + w / 2, t), "b": (l + w / 2, t + h)}[side]


dbox(s, "API\nPOST /conversations", **API, fill=TEAL)
dbox(s, "Direct upload\n(no API call)", **DIRECT, fill=AMBER)
dbox(s, "S3\nincoming/{id}.json", **S3, fill=SLATE)
dbox(s, "SQS\nprocessing queue", **SQS, fill=SLATE)
dbox(s, "Worker", **WORKER, fill=TEAL_DARK)
dbox(s, "OpenAI", **OPENAI, fill=NAVY)
dbox(s, "PostgreSQL", **PG, fill=NAVY)
dbox(s, "Review UI", **REVIEW, fill=SLATE)
dbox(s, "SQS DLQ", **DLQ, fill=DANGER)

# pending-row connector: API writes directly to Postgres, synchronously
a_top = edge(API, "t")
pg_top = edge(PG, "t")
elbow(s, a_top[0], a_top[1], 1.92, pg_top[0], pg_top[1], color=AMBER, dashed=False)
dlabel(s, (a_top[0] + pg_top[0]) / 2, 1.62, "pending row (sync write)", color=AMBER, size=10.5, width=3.4)

# main pipeline arrows
seg(s, *edge(API, "r"), *edge(S3, "l"), color=SLATE)
seg(s, *edge(DIRECT, "r"), *(edge(S3, "l")[0], edge(S3, "l")[1] + 0.001), color=SLATE)
seg(s, *edge(S3, "r"), *edge(SQS, "l"), color=SLATE)
dlabel(s, (edge(S3, "r")[0] + edge(SQS, "l")[0]) / 2, 2.5, "S3 event")
seg(s, *edge(SQS, "r"), *edge(WORKER, "l"), color=SLATE)
dlabel(s, (edge(SQS, "r")[0] + edge(WORKER, "l")[0]) / 2, 2.5, "poll")
seg(s, *edge(WORKER, "b"), *edge(OPENAI, "t"), color=TEAL_DARK)
dlabel(s, edge(WORKER, "b")[0] + 0.62, 3.62, "score", width=1.1)
seg(s, *edge(WORKER, "r"), *edge(PG, "l"), color=TEAL_DARK)
dlabel(s, (edge(WORKER, "r")[0] + edge(PG, "l")[0]) / 2, 2.5, "store result")
seg(s, *edge(PG, "b"), *edge(REVIEW, "t"), color=SLATE)
seg(s, *edge(SQS, "b"), *edge(DLQ, "t"), color=DANGER, dashed=True)
dlabel(s, edge(SQS, "b")[0] + 0.85, 3.78, "after 3 attempts", width=1.6, color=DANGER)

legend(s, [
    (SLATE, "pipeline flow", False),
    (AMBER, "pending-row write (sync)", False),
    (DANGER, "exhausted after retries", True),
], MARGIN, 4.98)

bullets(s, [
    "The API never calls OpenAI — it writes a pending row to Postgres, puts the object in S3, and returns 202",
    "A direct upload skips the API entirely — the same S3 → SQS → worker path picks it up",
    "One validation contract, one scoring path, one place the DLQ can be reached from",
], MARGIN, 5.4, 11.9, size=13.5, gap=8, height=1.3)

footer(s, "Architecture", 3)
add_notes(s, "The point of this diagram is 'one pipeline, two doors in.' Whether a conversation "
             "arrives via the API or lands in S3 directly, it converges on the exact same S3 -> "
             "SQS -> worker code path - there's no special-casing based on how it arrived. The "
             "amber line is the one part that's NOT on that shared path: the API writes the "
             "'pending' row straight to Postgres, synchronously, before it ever touches S3, so "
             "the id exists immediately and the caller gets a stable status URL right away.")

# ========================================================================
# Slide 4 - Ingestion paths
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Ingestion", "How conversations get in: two paths, one pipeline")

panel(s, MARGIN, 1.9, 5.55, 4.15, accent=TEAL)
panel_header(s, "Path 1 — API", MARGIN + 0.3, 2.1, 5.0, color=TEAL_DARK)
numbered_step(s, 1, "POST /conversations with the conversation text", MARGIN + 0.3, 2.68, 4.95, accent=TEAL)
numbered_step(s, 2, "Validated once; a “pending” row is written to Postgres immediately", MARGIN + 0.3, 3.68, 4.95, accent=TEAL)
numbered_step(s, 3, "Conversation is written to S3; API returns 202 without waiting for scoring", MARGIN + 0.3, 4.68, 4.95, accent=TEAL)

panel(s, MARGIN + 6.05, 1.9, 5.55, 4.15, accent=AMBER)
panel_header(s, "Path 2 — Direct S3 upload", MARGIN + 6.35, 2.1, 5.0, color=AMBER)
numbered_step(s, 1, "A file lands in S3 under incoming/ — no API call at all", MARGIN + 6.35, 2.68, 4.95, accent=AMBER)
numbered_step(s, 2, "An S3 event notification delivers straight to SQS", MARGIN + 6.35, 3.68, 4.95, accent=AMBER)
numbered_step(s, 3, "The worker creates the row itself, keyed by the S3 object key", MARGIN + 6.35, 4.68, 4.95, accent=AMBER)

banner = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MARGIN), Inches(6.18), Inches(11.9), Inches(0.5))
banner.fill.solid()
banner.fill.fore_color.rgb = NAVY
banner.line.fill.background()
no_shadow(banner)
rounded(banner, 0.25)
tf = banner.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Both converge on the same S3 → SQS → worker pipeline — one validation contract, one scoring code path"
_font(r, 13, WHITE, bold=True)

footer(s, "Ingestion paths", 4)
add_notes(s, "API path detail: the request body is validated through the same Pydantic model "
             "the worker also uses (rejects missing/empty/whitespace-only/non-string/over-20,000 "
             "character text). If the S3 write's outcome is genuinely unknown (e.g. a timeout), "
             "the row stays visible and pending rather than being silently deleted - deleting on "
             "an ambiguous failure could erase a conversation S3 actually accepted.\n\n"
             "Direct-upload detail: source_key (the S3 key) is the idempotency anchor, since a "
             "direct upload never gets an id from the API. A UNIQUE constraint on source_key "
             "means the same key uploaded twice finds the existing row instead of duplicating it.")

# ========================================================================
# Slide 5 - Scoring trigger, retries, DLQ
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Scoring", "How scoring is triggered — and what happens when it fails")

SQS2 = dict(left=0.9, top=2.55, w=1.8, h=0.85)
WK2 = dict(left=3.5, top=2.55, w=1.8, h=0.85)
PG2 = dict(left=8.9, top=2.55, w=1.9, h=0.85)
OA2 = dict(left=3.5, top=3.85, w=1.8, h=0.75)
DLQ2 = dict(left=0.9, top=3.85, w=1.8, h=0.75)

dbox(s, "SQS\nprocessing queue", **SQS2, fill=SLATE)
dbox(s, "Worker", **WK2, fill=TEAL_DARK)
dbox(s, "OpenAI", **OA2, fill=NAVY)
dbox(s, "PostgreSQL", **PG2, fill=NAVY)
dbox(s, "SQS DLQ", **DLQ2, fill=DANGER)

wk_top = edge(WK2, "t")
sqs_top = edge(SQS2, "t")
elbow(s, wk_top[0], wk_top[1], 2.25, sqs_top[0], sqs_top[1], color=SLATE, dashed=True)
dlabel(s, (wk_top[0] + sqs_top[0]) / 2, 1.95, "retry (visibility timeout)", color=SLATE, width=2.6)

seg(s, *edge(SQS2, "r"), *edge(WK2, "l"), color=SLATE)
dlabel(s, (edge(SQS2, "r")[0] + edge(WK2, "l")[0]) / 2, 2.22, "poll")
seg(s, *edge(WK2, "b"), *edge(OA2, "t"), color=TEAL_DARK)
dlabel(s, edge(WK2, "b")[0] + 0.6, 3.42, "score", width=1.0)
seg(s, *edge(WK2, "r"), *edge(PG2, "l"), color=TEAL_DARK)
dlabel(s, (edge(WK2, "r")[0] + edge(PG2, "l")[0]) / 2, 2.22, "commit result", width=3.6)
seg(s, *edge(SQS2, "b"), *edge(DLQ2, "t"), color=DANGER, dashed=True)
dlabel(s, edge(SQS2, "b")[0] + 0.75, 3.62, "after 3 attempts", width=1.5, color=DANGER)

bullets(s, [
    "The worker polls SQS continuously — a message stays invisible while it's being worked, and reappears automatically if not finished in time",
    "A transient failure is retried automatically; after 3 attempts SQS moves the message to the DLQ on its own",
    "Duplicate deliveries and out-of-order retries are both handled — the same object never scores into two rows, and a stale attempt can't overwrite a newer result",
], MARGIN, 4.95, 11.9, size=14, gap=11, height=1.8)

footer(s, "Scoring trigger", 5)
add_notes(s, "Trigger is SQS visibility, full stop - the worker is a simple poll loop, nothing "
             "depends on the API being up. Idempotency: source_key is UNIQUE. Claim fencing "
             "(the load-bearing safety net): every claim gets a fresh processing_token, so a "
             "superseded/slower attempt is rejected at the database level rather than silently "
             "overwriting a newer, correct result - without it, 'whoever writes last wins' could "
             "clobber a good row with a stale one. Honest remaining gap: a crash between "
             "'OpenAI answered' and 'the guarded commit landed' can still cause one extra paid "
             "OpenAI call - SQS is at-least-once and an LLM call can't be made atomic with a DB "
             "commit.")

# ========================================================================
# Slide 6 - OpenAI security and failure handling
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Scoring safely", "OpenAI: structured output, bounded retries, isolated secret")

bullets(s, [
    "The model returns a validated structured result — sentiment, a 0–1 risk score, a one-sentence rationale — not text parsed out of free-form output",
    "Requests are bounded by a timeout and a small number of automatic retries, so one slow or failing call can't stall the worker indefinitely",
    "A refusal is treated as a scoring failure like any other, and its explanation text is never stored or logged — it can echo back part of the flagged input",
    "The OpenAI API key is only ever injected into the worker — the API service never receives it, because the API never calls OpenAI",
], MARGIN, 1.95, 11.4, size=18, gap=20, height=4.6)

footer(s, "OpenAI integration", 6)
add_notes(s, "Every result also captures model name, prompt version, input/output token counts, "
             "and an estimated cost in USD from a hardcoded per-model price table (an estimate, "
             "not a billing feed). Retries use one mechanism (tenacity with exponential backoff) "
             "- the SDK's own built-in retries are explicitly disabled so there's exactly one "
             "retry policy to reason about, not two stacked on each other. The API-can't-leak-"
             "the-key point is a fact about the deployment, not just a policy: the key is never "
             "injected into that pod's environment at all.")

# ========================================================================
# Slide 7 - Durable storage & human review
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Storage & review", "PostgreSQL durability, and a human in the loop")

bullets(s, [
    "One table, one clear state machine — pending → processing → completed, or failed",
    "Postgres runs as a StatefulSet with a PersistentVolumeClaim — restart the pods and the data, including timestamps, survives unchanged",
    "Review is two server-rendered pages — a list and a detail view — no separate frontend to build or ship",
], MARGIN, 1.95, 6.5, size=16.5, gap=20, height=4.4)
add_picture_fit(s, SCREENSHOTS / "review-detail.png", 7.55, 1.9, 5.05, 4.65)

footer(s, "Storage and review", 7)
add_notes(s, "The screenshot is a real conversation from the deployed system, not a mock. "
             "Durability is demonstrated, not just claimed: the restart demo captures a "
             "completed row's full result before deleting the API, worker, and Postgres pods, "
             "then proves the row is byte-for-byte identical afterward, including "
             "completed_at - a re-scored duplicate would also say 'completed' but with a "
             "different timestamp, so that field is the real proof, not just the status.")

# ========================================================================
# Slide 8 - Observability & induced failure
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Observability", "Logs, metrics, a dashboard — and a failure induced on purpose")

bullets(s, [
    "Structured JSON logs and Prometheus metrics cover latency, processing outcomes, DLQ depth, and OpenAI tokens/cost",
    "A deterministic trigger fails one conversation on purpose — it retries 3 times and lands in the DLQ, while a healthy conversation submitted alongside it completes normally the whole time",
    "The dashboard and alerts show it live — DLQ depth rises during the failure, and the alert clears once the message is drained",
], MARGIN, 1.95, 6.5, size=15.5, gap=18, height=4.5)
add_picture_fit(s, SCREENSHOTS / "grafana-dashboard.png", 7.5, 1.85, 5.1, 3.55)
add_picture_fit(s, SCREENSHOTS / "prometheus-alerts.png", 7.5, 5.55, 5.1, 1.15)

footer(s, "Observability", 8)
add_notes(s, "On-call would watch: DLQ depth rising, two 'last progress' staleness gauges on the "
             "worker loop and DLQ collector (a bare TCP liveness check only proves the process "
             "is alive, not that it's doing anything - these catch a wedged loop that liveness "
             "alone would miss), the API's 5xx rate, and the cost counter's rate of change as a "
             "runaway-spend signal. Both screenshots are from the live deployed system. Worth "
             "mentioning if asked: found and fixed a real bug preparing these - the health "
             "gauges are defined in a module shared by both API and worker processes, so "
             "without a job-label filter the API's own always-zero copy made both alerts fire "
             "permanently. Exactly the kind of false positive that trains on-call to ignore "
             "alerts.")

# ========================================================================
# Slide 9 - Infrastructure
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Infrastructure", "Terraform, one image, Helm, and secrets kept out of the chart")

bullets(s, [
    "Terraform provisions the S3 bucket, the processing queue, its DLQ, and the event notification wiring them together",
    "One Docker image, two runtime commands — the API and worker run the same image; Helm just overrides the container command",
    "Deployed via Helm: non-root, read-only filesystem, health probes tuned to real startup behavior",
    "Secrets are created out-of-band and referenced by name — never templated into the chart or committed to Git; IAM policies follow least privilege per component",
], MARGIN, 1.95, 11.6, size=16.5, gap=17, height=4.6)

footer(s, "Infrastructure", 9)
add_notes(s, "Docker: multi-stage build - a builder stage installs dependencies into a venv, the "
             "runtime stage copies only that venv plus app/, no git metadata/tests/Terraform "
             "state. Deploy safety: deploy.sh tags images with the current git SHA and refuses "
             "to build at all if any tracked source/Docker/Helm/Terraform file has uncommitted "
             "changes, so one tag can never silently point at two different images. Least "
             "privilege specifics: the API can only PutObject under incoming/*; the worker can "
             "GetObject there and Receive/Delete/GetAttributes only on the processing queue plus "
             "GetUrl/GetAttributes on the DLQ - never a wildcard resource. Honest caveat: "
             "LocalStack Community doesn't enforce IAM at the API-call level, so these policies "
             "show the intended real-AWS shape, not a currently-enforced restriction.")

# ========================================================================
# Slide 10 - Production deployment roadmap
# ========================================================================
s = add_slide(); set_background(s)
kicker_title(s, "Looking ahead", "Production deployment roadmap")

cap = s.shapes.add_textbox(Inches(MARGIN), Inches(1.58), Inches(11.6), Inches(0.3))
p = cap.text_frame.paragraphs[0]
r = p.add_run()
r.text = "A proposed CI/CD path for production — not part of this take-home's current deployment"
_font(r, 12, MUTED, italic=True)

RB_W, RB_H, RB_GAP = 1.75, 0.95, 0.22
RB_Y = 2.05
rb_x = [MARGIN + i * (RB_W + RB_GAP) for i in range(6)]
REPO = dict(left=rb_x[0], top=RB_Y, w=RB_W, h=RB_H)
CI = dict(left=rb_x[1], top=RB_Y, w=RB_W, h=RB_H)
REG = dict(left=rb_x[2], top=RB_Y, w=RB_W, h=RB_H)
GITOPS = dict(left=rb_x[3], top=RB_Y, w=RB_W, h=RB_H)
ARGO = dict(left=rb_x[4], top=RB_Y, w=RB_W, h=RB_H)
PROD = dict(left=rb_x[5], top=RB_Y, w=RB_W, h=RB_H)

dbox(s, "Application\nrepository", **REPO, fill=SLATE)
dbox(s, "CI pipeline", **CI, fill=TEAL)
dbox(s, "Container\nregistry", **REG, fill=SLATE)
dbox(s, "GitOps\nrepository", **GITOPS, fill=AMBER)
dbox(s, "Argo CD", **ARGO, fill=TEAL_DARK)
dbox(s, "Kubernetes\nproduction\ncluster", **PROD, fill=NAVY, size=11.5)

stages = [REPO, CI, REG, GITOPS, ARGO, PROD]
arrow_labels = ["commit", "build & push", "update tag", "git sync", "deploy"]
for a, b, lbl in zip(stages, stages[1:], arrow_labels):
    seg(s, *edge(a, "r"), *edge(b, "l"), color=SLATE)
    dlabel(s, (edge(a, "r")[0] + edge(b, "l")[0]) / 2, RB_Y + RB_H + 0.1, lbl, width=RB_W)

CARD_W, CARD_H, CARD_GAP_X, CARD_GAP_Y = 5.55, 1.35, 0.5, 0.2
CARD_Y1 = 3.65
CARD_Y2 = CARD_Y1 + CARD_H + CARD_GAP_Y
CARD_X1 = MARGIN
CARD_X2 = MARGIN + CARD_W + CARD_GAP_X

roadmap_points = [
    (SLATE, "1", "Repository separation",
     "App code and the Dockerfile stay here; Helm config moves to a separate GitOps repo — Terraform can split out too.",
     CARD_X1, CARD_Y1),
    (TEAL, "2", "CI pipeline",
     "Every change: tests, security scan, build, tag with the commit SHA, push to the registry, then update the GitOps repo's tag.",
     CARD_X2, CARD_Y1),
    (TEAL_DARK, "3", "Argo CD deployment",
     "Argo CD syncs the GitOps repo into the cluster — changes reviewed via Git, rollback by reverting to a known-good commit.",
     CARD_X1, CARD_Y2),
    (NAVY, "4", "Production reliability",
     "Managed RDS/SQS/S3/DLQ, a cloud secret manager and workload identity, multiple replicas, queue-depth autoscaling, monitoring and backups.",
     CARD_X2, CARD_Y2),
]
for accent, num, head, body, cx, cy in roadmap_points:
    panel(s, cx, cy, CARD_W, CARD_H, accent=accent)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.25), Inches(cy + 0.2), Inches(0.32), Inches(0.32))
    badge.fill.solid()
    badge.fill.fore_color.rgb = accent
    badge.line.fill.background()
    no_shadow(badge)
    btf = badge.text_frame
    btf.word_wrap = False
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    btf.margin_left = 0
    btf.margin_right = 0
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run()
    br.text = num
    _font(br, 12.5, WHITE, bold=True)

    ht = s.shapes.add_textbox(Inches(cx + 0.7), Inches(cy + 0.15), Inches(CARD_W - 1.0), Inches(0.35))
    hp = ht.text_frame.paragraphs[0]
    hr = hp.add_run()
    hr.text = head
    _font(hr, 14.5, NAVY, bold=True)

    dt = s.shapes.add_textbox(Inches(cx + 0.25), Inches(cy + 0.58), Inches(CARD_W - 0.5), Inches(0.7))
    dtf = dt.text_frame
    dtf.word_wrap = True
    dp = dtf.paragraphs[0]
    dp.line_spacing = 1.05
    dr = dp.add_run()
    dr.text = body
    _font(dr, 12.5, INK)

footer(s, "Production roadmap", 10)
add_notes(s, "This is a proposed future CI/CD path, not something implemented in this "
             "take-home - worth being explicit about that distinction if asked. The flow: a "
             "commit to the application repo triggers CI (tests, security scanning, image "
             "build, SHA tag, push to a registry), which then updates the image tag in a "
             "separate GitOps repo; Argo CD watches that repo and reconciles the cluster to "
             "match it. Rollback becomes 'revert a Git commit,' not 'run a manual kubectl "
             "command.' Separating the GitOps repo from the application repo also means "
             "environment-specific Helm values (dev/staging/prod) never need to touch the "
             "application codebase.\n\n"
             "Honest tradeoffs/exclusions worth naming if asked, since they're not their own "
             "slide anymore: no exactly-once processing claim anywhere (SQS is at-least-once "
             "by design); a crash between an OpenAI response and the DB commit can still cause "
             "one extra paid call, bounded and rare, not eliminated (narrowed - not solved - by "
             "an idempotency key on the LLM call, not a transactional outbox, since the "
             "non-atomic boundary is an external API call, not two local writes); LocalStack "
             "doesn't enforce the IAM policies Terraform defines; and the take-home was scoped "
             "down on purpose - no CI, no service mesh, no ORM, no new database.")

prs.save(str(OUT))
print(f"Wrote {OUT} ({len(prs.slides)} slides)")
